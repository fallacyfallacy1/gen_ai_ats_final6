import google.generativeai as genai
import streamlit as st
from fpdf import FPDF
from docx import Document
from io import BytesIO
import re
from sklearn.feature_extraction.text import CountVectorizer
from PyPDF2 import PdfReader

# Configure Generative AI
genai.configure(api_key="AIzaSyCmDooErfdJkWyrPAecchTor3j30v3tcY8")
model = genai.GenerativeModel("gemini-1.5-flash")

# Function to interact with Generative AI
def ask_gemini(prompt):
    response = model.generate_content(prompt)
    return response.text

# Function to extract text from a PDF file
def extract_text_from_pdf(pdf_file):
    pdf_reader = PdfReader(pdf_file)
    text = ""
    for page in pdf_reader.pages:
        text += page.extract_text()
    return text

# Function to calculate ATS score
def calculate_ats_score(resume, job_description):
    def clean_text(text):
        return re.sub(r'[^\w\s]', '', text.lower())

    resume_clean = clean_text(resume)
    job_description_clean = clean_text(job_description)

    vectorizer = CountVectorizer(max_features=20, stop_words='english')
    vectorizer.fit([job_description_clean])
    keywords = vectorizer.get_feature_names_out()

    matched_keywords = [word for word in keywords if word in resume_clean]
    keyword_match_score = len(matched_keywords) / len(keywords) * 100

    return keyword_match_score

# Function to create a PDF resume with visual elements
def create_resume_pdf(resume_text):
    pdf = FPDF()
    pdf.add_page()

    # Add DejaVu font
    pdf.add_font('DejaVu', '', r'C:\Users\falla\my_env\DejaVuSansCondensed.ttf', uni=True)
    pdf.set_font('DejaVu', '', 16)

    # Add header
    pdf.set_fill_color(0, 102, 204)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(0, 10, "Resume", ln=True, align="C", fill=True)

    # Add body text
    pdf.set_text_color(0, 0, 0)
    pdf.set_font('DejaVu', '', 12)
    pdf.multi_cell(0, 10, resume_text)

    return pdf.output(dest='S').encode('latin-1')

# Function to populate a Word resume template
def populate_template(template_path, placeholders):
    doc = Document(template_path)
    for paragraph in doc.paragraphs:
        for placeholder, value in placeholders.items():
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, value)

    # Save populated resume to bytes
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream.getvalue()

# Streamlit UI
st.title("Resume Builder and ATS Checker")

# Resume modification and job description inputs
instr = st.text_input("Modify the resume structure/instructions:")
jd = st.text_area("Job description:")

# Resume upload and input
uploaded_file = st.file_uploader("Upload PDF Resume", type=["pdf"])
resume_text = st.text_area("Or paste your resume text:")

if uploaded_file is not None:
    resume_text = extract_text_from_pdf(uploaded_file)
    st.success("Resume text extracted from PDF!")

# Prompt for Generative AI
prompt = (
    f"As an experienced career consultant (HR), generate an appropriate resume "
    f"based on the following instructions: {instr} and job description: {jd}.\n\n"
    f"Current Resume:\n{resume_text}"
    f"Write in HTML format"
)

if st.button("Generate Resume!"):
    if resume_text and jd:
        response = ask_gemini(prompt)
        st.write("Generated Resume:")
        st.write(response)

        # Generate PDF and Word resumes
        pdf_bytes = create_resume_pdf(response)

        placeholders = {"{{ResumeContent}}": response}
        word_template_path = r"C:\Users\falla\my_env\resumeformat.docx"  # Replace with your downloaded template path
        docx_bytes = populate_template(word_template_path, placeholders)

        # Provide download buttons
        st.download_button(
            label="Download Resume as PDF",
            data=pdf_bytes,
            file_name="generated_resume.pdf",
            mime="application/pdf"
        )
        st.download_button(
            label="Download Resume as Word Document",
            data=docx_bytes,
            file_name="generated_resume.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    else:
        st.error("Please provide both a resume and a job description.")

# ATS Score Checker
ats_resume = resume_text
ats_job_description = st.text_area("Paste the job description (for ATS Score):")

if st.button("Check ATS Score"):
    if ats_resume and ats_job_description:
        score = calculate_ats_score(ats_resume, ats_job_description)
        st.success(f"Your ATS Score is: {score:.2f}%")
    else:
        st.error("Please provide both a resume and a job description.")


